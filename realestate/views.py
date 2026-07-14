from django.shortcuts import render, redirect, get_object_or_404
from django.db.models import Sum, Q
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.http import HttpResponseForbidden
import io
from django.utils import timezone
from datetime import timedelta, date
from .models import Landlord, Property, Unit, Contract, Collection, Tenant, DocumentFolder, DocumentFile, \
    ApprovalDocument, PropertyDocument, SyncLog

hidden_accounts = ['Landlord (ID: 1)', 'YBN General Portfolio', 'Accounts', 'Output VAT on Sale']


def get_dynamic_notifications(landlord):
    notifs = []

    # Extract the two most recent syncs to establish the delta window
    syncs = list(SyncLog.objects.order_by('-timestamp')[:2])

    if len(syncs) >= 2:
        previous_sync = syncs[1].timestamp
        latest_sync = syncs[0].timestamp
        sync_display = latest_sync.astimezone().strftime("%I:%M %p")
    elif len(syncs) == 1:
        previous_sync = timezone.now() - timedelta(days=1)
        latest_sync = syncs[0].timestamp
        sync_display = latest_sync.astimezone().strftime("%I:%M %p")
    else:
        previous_sync = timezone.now() - timedelta(days=1)
        sync_display = "Recently"

    # 1. Collection Status Changes Since Last Sync
    recent_changes = Collection.objects.filter(
        contract__unit__property_asset__owner=landlord,
        last_updated__gt=previous_sync
    )

    stats = recent_changes.order_by().values('status').annotate(total=Sum('amount'))

    for stat in stats:
        status = stat['status']
        total = stat['total'] or 0
        if total > 0:
            icon = "fa-bell"
            color = "var(--primary-gold)"
            if status == "Cleared":
                icon, color = "fa-check-double", "var(--success)"
            elif status == "Bounced":
                icon, color = "fa-triangle-exclamation", "var(--danger)"
            elif status == "Pending PDC":
                icon, color = "fa-clock", "var(--warning)"
            elif status == "Terminated PDC":
                icon, color = "fa-ban", "#64748b"
            elif "Replaced" in status:
                icon, color = "fa-rotate", "#3b82f6"

            notifs.append({
                'icon': icon, 'color': color,
                'title': f'{status} Updates',
                'desc': f'AED {total:,.2f} added to {status} (last sync: {sync_display}).'
            })

    # 2. Fetch Pending Approvals
    pending_approvals = ApprovalDocument.objects.filter(
        Q(landlord=landlord) | Q(landlord__isnull=True),
        status='UnApproved'
    ).count()

    if pending_approvals > 0:
        notifs.append({
            'icon': 'fa-file-signature', 'color': 'var(--warning)',
            'title': 'Approvals Required',
            'desc': f"You have {pending_approvals} UnApproved vouchers pending."
        })

    return notifs


# ==========================================
# PUBLIC & AUTH ROUTES
# ==========================================

def landing_page(request):
    if request.user.is_authenticated:
        if request.user.is_superuser:
            return redirect('admin_dashboard')
        if hasattr(request.user, 'landlord_profile'):
            return redirect('owner_dashboard', landlord_id=request.user.landlord_profile.id)
    return render(request, 'realestate/landing_page.html')


def login_view(request):
    if request.user.is_authenticated:
        if request.user.is_superuser:
            return redirect('admin_dashboard')
        if hasattr(request.user, 'landlord_profile'):
            return redirect('owner_dashboard', landlord_id=request.user.landlord_profile.id)

    role = request.GET.get('role', 'Executive')

    if request.method == 'POST':
        u = request.POST.get('username', '').strip()
        p = request.POST.get('password', '').strip()

        # Explicit backend check for empty fields
        if not u or not p:
            return render(request, 'realestate/login.html', {
                'error': 'Please fill out both username and password fields.',
                'role': role
            })

        user = authenticate(request, username=u, password=p)
        if user is not None:
            login(request, user)
            if user.is_superuser:
                return redirect('admin_dashboard')
            if hasattr(user, 'landlord_profile'):
                return redirect('owner_dashboard', landlord_id=user.landlord_profile.id)
            return redirect('landing_page')
        else:
            # Custom error for wrong credentials
            return render(request, 'realestate/login.html', {
                'error': 'Incorrect username or password. Please try again.',
                'role': role
            })

    return render(request, 'realestate/login.html', {'role': role})

def logout_view(request):
    logout(request)
    return redirect('landing_page')


# ==========================================
# ADMIN ROUTES
# ==========================================

@login_required(login_url='/login/')
def admin_dashboard(request):
    if not request.user.is_superuser:
        return HttpResponseForbidden("<h1>403 - Security Violation</h1><p>Admin clearance required.</p>")
    landlords = Landlord.objects.exclude(name__in=hidden_accounts).order_by('name')
    return render(request, 'realestate/portal_home.html', {'landlords': landlords})


# ==========================================
# OWNER ROUTES
# ==========================================

@login_required(login_url='/login/')
def owner_dashboard(request, landlord_id):
    landlord = get_object_or_404(Landlord, id=landlord_id)
    if not request.user.is_superuser:
        if not hasattr(request.user, 'landlord_profile') or request.user.landlord_profile.id != landlord.id:
            return HttpResponseForbidden("<h1>403 - Security Violation</h1>")

    properties = Property.objects.filter(owner=landlord)
    total_units = Unit.objects.filter(property_asset__in=properties).count()
    occupied_units = Unit.objects.filter(property_asset__in=properties, contracts__status='Active').distinct().count()
    occupancy_rate = int((occupied_units / total_units) * 100) if total_units > 0 else 0

    today = timezone.now().date()
    first_of_month = date(today.year, today.month, 1)

    kpi_collections = Collection.objects.filter(
        contract__unit__property_asset__in=properties,
        due_date__gte=first_of_month,
        due_date__lte=today
    )

    kpi_breakdown = kpi_collections.order_by().values('status').annotate(total=Sum('amount'))
    kpi_dict = {item['status']: item['total'] or 0 for item in kpi_breakdown}

    kpi_sums = {
        'cleared': kpi_dict.get('Cleared', 0),
        'pending': kpi_dict.get('Pending PDC', 0),
        'bounced': kpi_dict.get('Bounced', 0),
        'replaced': kpi_dict.get('Replaced', 0),
        'partial': kpi_dict.get('Replaced/Partial', 0),
        'terminated': kpi_dict.get('Terminated PDC', 0),
    }

    recent_transactions = Collection.objects.filter(
        contract__unit__property_asset__in=properties
    ).order_by('-id')[:20]

    notifications = get_dynamic_notifications(landlord)

    context = {
        'landlord': landlord,
        'properties_count': properties.count(),
        'total_units': total_units,
        'occupancy_rate': occupancy_rate,
        'kpi_sums': kpi_sums,
        'recent_transactions': recent_transactions,
        'notifications': notifications,
        'notif_count': len(notifications)
    }
    return render(request, 'realestate/owner_dashboard.html', context)


@login_required(login_url='/login/')
def owner_properties(request, landlord_id):
    landlord = get_object_or_404(Landlord, id=landlord_id)

    if not request.user.is_superuser:
        if not hasattr(request.user, 'landlord_profile') or request.user.landlord_profile.id != landlord.id:
            return HttpResponseForbidden("<h1>403 - Security Violation</h1>")

    today = timezone.now().date()
    thirty_days = today + timedelta(days=30)
    ninety_days = today + timedelta(days=90)

    active_contracts = Contract.objects.filter(
        unit__property_asset__owner=landlord,
        status='Active'
    ).select_related('unit__property_asset', 'tenant')

    renewals = active_contracts.filter(end_date__lte=thirty_days).order_by('-end_date')
    expiry_90 = active_contracts.filter(end_date__gt=thirty_days, end_date__lte=ninety_days).order_by('end_date')

    vacant_units = Unit.objects.filter(
        property_asset__owner=landlord
    ).exclude(
        contracts__status='Active'
    ).select_related('property_asset').order_by('property_asset__name', 'unit_number')

    notifications = get_dynamic_notifications(landlord)

    context = {
        'landlord': landlord,
        'active_contracts': active_contracts.order_by('-start_date'),
        'renewals': renewals,
        'expiry_90': expiry_90,
        'vacant_units': vacant_units,
        'today': today,
        'notifications': notifications,
        'notif_count': len(notifications)
    }
    return render(request, 'realestate/owner_properties.html', context)


@login_required(login_url='/login/')
def owner_vault(request, landlord_id):
    landlord = get_object_or_404(Landlord, id=landlord_id)

    if not request.user.is_superuser:
        if not hasattr(request.user, 'landlord_profile') or request.user.landlord_profile.id != landlord.id:
            return HttpResponseForbidden("<h1>403 - Security Violation</h1>")

    transactions = Collection.objects.filter(
        contract__unit__property_asset__owner=landlord
    ).select_related(
        'contract__tenant',
        'contract__unit__property_asset'
    )

    properties = Property.objects.filter(owner=landlord)

    search_query = request.GET.get('search', '')
    property_filter = request.GET.get('property', '')
    status_filter = request.GET.get('status', '')
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    sort_by = request.GET.get('sort', '-id')

    if search_query:
        transactions = transactions.filter(
            Q(contract__tenant__full_name__icontains=search_query) |
            Q(cheque_number__icontains=search_query)
        )
    if property_filter:
        transactions = transactions.filter(contract__unit__property_asset_id=property_filter)
    if status_filter:
        transactions = transactions.filter(status=status_filter)
    if date_from:
        transactions = transactions.filter(due_date__gte=date_from)
    if date_to:
        transactions = transactions.filter(due_date__lte=date_to)

    if sort_by == 'date_asc':
        transactions = transactions.order_by('due_date')
    elif sort_by == '-due_date':
        transactions = transactions.order_by('-due_date')
    elif sort_by == 'amount_desc':
        transactions = transactions.order_by('-amount')
    elif sort_by == 'amount_asc':
        transactions = transactions.order_by('amount')
    elif sort_by == 'id_asc':
        transactions = transactions.order_by('id')
    else:
        transactions = transactions.order_by('-id')

    vault_breakdown = transactions.order_by().values('status').annotate(total=Sum('amount'))
    vault_dict = {item['status']: item['total'] or 0 for item in vault_breakdown}

    vault_sums = {
        'cleared': vault_dict.get('Cleared', 0),
        'pending': vault_dict.get('Pending PDC', 0),
        'bounced': vault_dict.get('Bounced', 0),
        'replaced': vault_dict.get('Replaced', 0),
        'partial': vault_dict.get('Replaced/Partial', 0),
        'terminated': vault_dict.get('Terminated PDC', 0),
    }

    notifications = get_dynamic_notifications(landlord)

    context = {
        'landlord': landlord,
        'transactions': transactions,
        'properties': properties,
        'vault_sums': vault_sums,
        'filters': {
            'search': search_query,
            'property': int(property_filter) if property_filter.isdigit() else '',
            'status': status_filter,
            'date_from': date_from,
            'date_to': date_to,
            'sort': sort_by,
        },
        'notifications': notifications,
        'notif_count': len(notifications)
    }
    return render(request, 'realestate/owner_vault.html', context)


@login_required(login_url='/login/')
def owner_view_properties(request, landlord_id):
    landlord = get_object_or_404(Landlord, id=landlord_id)
    if not request.user.is_superuser:
        if not hasattr(request.user, 'landlord_profile') or request.user.landlord_profile.id != landlord.id:
            return HttpResponseForbidden("<h1>403 - Security Violation</h1>")

    properties = Property.objects.filter(owner=landlord)
    property_data = []

    for p in properties:
        total_units = Unit.objects.filter(property_asset=p).count()
        occupied_units = Unit.objects.filter(property_asset=p, contracts__status='Active').distinct().count()
        occupancy_rate = int((occupied_units / total_units) * 100) if total_units > 0 else 0

        active_contracts = Contract.objects.filter(unit__property_asset=p, status='Active').select_related('tenant',
                                                                                                           'unit')
        documents = p.documents.all().order_by('-uploaded_at')

        property_data.append({
            'property': p,
            'occupancy_rate': occupancy_rate,
            'active_contracts': active_contracts,
            'documents': documents,
        })

    notifications = get_dynamic_notifications(landlord)

    return render(request, 'realestate/view_properties.html', {
        'landlord': landlord,
        'property_data': property_data,
        'notifications': notifications,
        'notif_count': len(notifications)
    })


import csv
from django.http import HttpResponse


# --- 1. UPDATE YOUR EXISTING TENANT DIRECTORY VIEW ---
@login_required(login_url='/login/')
def owner_tenant_details(request, landlord_id):
    landlord = get_object_or_404(Landlord, id=landlord_id)
    if not request.user.is_superuser:
        if not hasattr(request.user, 'landlord_profile') or request.user.landlord_profile.id != landlord.id:
            return HttpResponseForbidden("<h1>403</h1>")

    search_query = request.GET.get('search', '').strip()
    prop_filter = request.GET.get('property', '')
    unit_filter = request.GET.get('unit', '').strip()
    lease_status = request.GET.get('lease_status', '')  # NEW FILTER

    properties = Property.objects.filter(owner=landlord)
    tenants_query = Tenant.objects.filter(contracts__unit__property_asset__owner=landlord)

    if search_query:
        tenants_query = tenants_query.filter(
            Q(full_name__icontains=search_query) |
            Q(phone__icontains=search_query) |
            Q(emirates_id__icontains=search_query)
        )
    if prop_filter:
        tenants_query = tenants_query.filter(contracts__unit__property_asset_id=prop_filter)
    if unit_filter:
        tenants_query = tenants_query.filter(contracts__unit__unit_number__icontains=unit_filter)
    if lease_status:
        tenants_query = tenants_query.filter(contracts__status=lease_status)

    tenants = tenants_query.distinct()
    tenant_data = []

    for t in tenants:
        contracts = Contract.objects.filter(tenant=t, unit__property_asset__owner=landlord).select_related(
            'unit__property_asset')
        collections = Collection.objects.filter(contract__in=contracts).order_by('-due_date')

        cleared = collections.filter(status='Cleared').aggregate(Sum('amount'))['amount__sum'] or 0
        pending = collections.filter(status='Pending PDC').aggregate(Sum('amount'))['amount__sum'] or 0
        bounced = collections.filter(status='Bounced').aggregate(Sum('amount'))['amount__sum'] or 0

        tenant_data.append({
            'tenant': t,
            'contracts': contracts,
            'financials': {'cleared': cleared, 'pending': pending, 'bounced': bounced}
        })

    notifications = get_dynamic_notifications(landlord)

    return render(request, 'realestate/tenant_details.html', {
        'landlord': landlord,
        'tenant_data': tenant_data,
        'properties': properties,
        'filters': {
            'search': search_query,
            'property': int(prop_filter) if prop_filter.isdigit() else '',
            'unit': unit_filter,
            'lease_status': lease_status  # ADDED TO CONTEXT
        },
        'notifications': notifications,
        'notif_count': len(notifications)
    })


# --- 2. ADD THE NEW REPORTS VIEW ---
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    Presentation = None


# --- REPORT VIEW WITH COMPREHENSIVE PPTX ENGINE ---
@login_required(login_url='/login/')
def owner_reports(request, landlord_id):
    landlord = get_object_or_404(Landlord, id=landlord_id)
    if not request.user.is_superuser:
        if not hasattr(request.user, 'landlord_profile') or request.user.landlord_profile.id != landlord.id:
            return HttpResponseForbidden("<h1>403</h1>")

    prop_filter = request.GET.get('property', '')
    unit_filter = request.GET.get('unit', '').strip()
    tenant_filter = request.GET.get('tenant', '').strip()
    status_filter = request.GET.get('status', '')
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    export_format = request.GET.get('export', '')

    # Base query for payments
    reports_data = Collection.objects.filter(contract__unit__property_asset__owner=landlord).select_related(
        'contract__tenant', 'contract__unit__property_asset'
    ).order_by('-due_date')

    # Apply filters
    if prop_filter:
        reports_data = reports_data.filter(contract__unit__property_asset_id=prop_filter)
    if unit_filter:
        reports_data = reports_data.filter(contract__unit__unit_number__icontains=unit_filter)
    if tenant_filter:
        reports_data = reports_data.filter(contract__tenant__full_name__icontains=tenant_filter)
    if status_filter:
        reports_data = reports_data.filter(status=status_filter)
    if date_from:
        reports_data = reports_data.filter(due_date__gte=date_from)
    if date_to:
        reports_data = reports_data.filter(due_date__lte=date_to)

    # -------------------------------------------------------------
    # PPTX EXPORT ENGINE
    # -------------------------------------------------------------
    if export_format == 'ppt':
        if Presentation is None:
            return HttpResponse("<h1>Please run 'pip install python-pptx' on your server to use this feature.</h1>")

        # Initialize widescreen presentation (16:9 ratio)
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]  # Completely blank layout to build custom UI

        # Brand Color Palette
        DARK_BG = RGBColor(17, 24, 39)  # #111827 Charcoal
        GOLD = RGBColor(212, 175, 55)  # #D4AF37 Gold
        WHITE = RGBColor(255, 255, 255)
        TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B Slate Gray
        TEXT_MAIN = RGBColor(30, 41, 59)  # #1E293B Dark Blue/Gray

        # Calculate dynamic metrics for slides based on current filtered subset
        cleared_total = reports_data.filter(status='Cleared').aggregate(Sum('amount'))['amount__sum'] or 0
        pending_total = reports_data.filter(status='Pending PDC').aggregate(Sum('amount'))['amount__sum'] or 0
        bounced_total = reports_data.filter(status='Bounced').aggregate(Sum('amount'))['amount__sum'] or 0
        grand_total = reports_data.aggregate(Sum('amount'))['amount__sum'] or 0

        # -------------------------------------------------------------
        # SLIDE 1: Title Slide (Dark Brand Theme)
        # -------------------------------------------------------------
        slide1 = prs.slides.add_slide(blank_layout)

        # Add full dark background shape
        bg_shape = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = DARK_BG
        bg_shape.line.color.rgb = DARK_BG

        # Title text box
        title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.5))
        tf = title_box.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = "PORTFOLIO PERFORMANCE REPORT"
        p1.font.name = "Arial"
        p1.font.size = Pt(44)
        p1.font.bold = True
        p1.font.color.rgb = GOLD
        p1.alignment = PP_ALIGN.LEFT

        p2 = tf.add_paragraph()
        p2.text = f"Prepared for: {landlord.name}"
        p2.font.name = "Arial"
        p2.font.size = Pt(20)
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(20)

        p3 = tf.add_paragraph()
        date_str = timezone.now().strftime("%B %d, %Y")
        p3.text = f"Generated on: {date_str}  |  YBN Portal Engine"
        p3.font.name = "Arial"
        p3.font.size = Pt(13)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(10)

        # -------------------------------------------------------------
        # SLIDE 2: Executive Financial Summary (Light Theme)
        # -------------------------------------------------------------
        slide2 = prs.slides.add_slide(blank_layout)

        # Header text
        header_box = slide2.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.833), Inches(1.0))
        tf_header = header_box.text_frame
        p_head = tf_header.paragraphs[0]
        p_head.text = "EXECUTIVE FINANCIAL SUMMARY"
        p_head.font.name = "Arial"
        p_head.font.size = Pt(26)
        p_head.font.bold = True
        p_head.font.color.rgb = DARK_BG

        # Helper card generator function (creates beautiful metric cards side-by-side)
        def create_metric_card(slide, left, top, width, height, title_text, amount_value, accent_rgb):
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
            card.fill.solid()
            card.fill.fore_color.rgb = WHITE
            card.line.color.rgb = accent_rgb
            card.line.width = Pt(2)

            ctf = card.text_frame
            ctf.word_wrap = True

            p_title = ctf.paragraphs[0]
            p_title.text = title_text.upper()
            p_title.font.name = "Arial"
            p_title.font.size = Pt(11)
            p_title.font.bold = True
            p_title.font.color.rgb = TEXT_MUTED
            p_title.alignment = PP_ALIGN.CENTER

            p_amt = ctf.add_paragraph()
            p_amt.text = f"AED {amount_value:,.2f}"
            p_amt.font.name = "Arial"
            p_amt.font.size = Pt(22)
            p_amt.font.bold = True
            p_amt.font.color.rgb = DARK_BG
            p_amt.space_before = Pt(12)
            p_amt.alignment = PP_ALIGN.CENTER

        # Draw 3 metric cards cleanly aligned
        card_width = Inches(3.6)
        card_height = Inches(1.8)
        top_y = Inches(1.8)

        create_metric_card(slide2, Inches(0.75), top_y, card_width, card_height, "Cleared Collections", cleared_total,
                           RGBColor(22, 163, 74))  # Green
        create_metric_card(slide2, Inches(4.85), top_y, card_width, card_height, "Pending PDCs", pending_total,
                           RGBColor(217, 119, 6))  # Amber
        create_metric_card(slide2, Inches(8.95), top_y, card_width, card_height, "Bounced Cheques", bounced_total,
                           RGBColor(220, 38, 38))  # Red

        # Add Large Grand Total Banner below cards
        banner = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(4.1), Inches(11.83), Inches(1.4))
        banner.fill.solid()
        banner.fill.fore_color.rgb = DARK_BG
        banner.line.color.rgb = GOLD
        banner.line.width = Pt(1.5)

        btf = banner.text_frame
        bp1 = btf.paragraphs[0]
        bp1.text = "COMBINED PORTFOLIO REVENUE UNDER CONTRACT"
        bp1.font.name = "Arial"
        bp1.font.size = Pt(10)
        bp1.font.bold = True
        bp1.font.color.rgb = GOLD
        bp1.alignment = PP_ALIGN.CENTER

        bp2 = btf.add_paragraph()
        bp2.text = f"AED {grand_total:,.2f}"
        bp2.font.name = "Arial"
        bp2.font.size = Pt(32)
        bp2.font.bold = True
        bp2.font.color.rgb = WHITE
        bp2.alignment = PP_ALIGN.CENTER
        bp2.space_before = Pt(8)

        # Footer
        footer_box = slide2.shapes.add_textbox(Inches(0.75), Inches(6.8), Inches(11.83), Inches(0.4))
        ftf = footer_box.text_frame
        fp = ftf.paragraphs[0]
        fp.text = "YBN Real Estate Portal  |  Financial Performance Dashboard Summary"
        fp.font.name = "Arial"
        fp.font.size = Pt(9)
        fp.font.color.rgb = TEXT_MUTED

        # -------------------------------------------------------------
        # SLIDE 3: Property Breakdown Table
        # -------------------------------------------------------------
        slide3 = prs.slides.add_slide(blank_layout)

        # Slide Title
        header_box3 = slide3.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.833), Inches(1.0))
        tf_header3 = header_box3.text_frame
        p_head3 = tf_header3.paragraphs[0]
        p_head3.text = "PORTFOLIO PROPERTY BREAKDOWN"
        p_head3.font.name = "Arial"
        p_head3.font.size = Pt(26)
        p_head3.font.bold = True
        p_head3.font.color.rgb = DARK_BG

        # Aggregate data by property
        properties = Property.objects.filter(owner=landlord)
        property_summaries = []
        for prop in properties:
            prop_data = reports_data.filter(contract__unit__property_asset=prop)
            prop_total = prop_data.aggregate(Sum('amount'))['amount__sum'] or 0
            prop_bounced = prop_data.filter(status='Bounced').aggregate(Sum('amount'))['amount__sum'] or 0
            prop_pending = prop_data.filter(status='Pending PDC').aggregate(Sum('amount'))['amount__sum'] or 0
            prop_cleared = prop_data.filter(status='Cleared').aggregate(Sum('amount'))['amount__sum'] or 0

            if prop_total > 0:  # Only show active properties in slide table
                property_summaries.append({
                    'name': prop.name,
                    'cleared': prop_cleared,
                    'pending': prop_pending,
                    'bounced': prop_bounced,
                    'total': prop_total
                })

        # Render PowerPoint Table
        rows = len(property_summaries) + 1  # data rows + header
        cols = 5
        left = Inches(0.75)
        top = Inches(1.6)
        width = Inches(11.83)
        height = Inches(0.5 * rows)

        table_shape = slide3.shapes.add_table(rows, cols, left, top, width, height)
        table = table_shape.table

        # Column widths
        table.columns[0].width = Inches(3.83)
        table.columns[1].width = Inches(2.0)
        table.columns[2].width = Inches(2.0)
        table.columns[3].width = Inches(2.0)
        table.columns[4].width = Inches(2.0)

        # Style Headers
        headers = ["PROPERTY ASSET", "CLEARED AMOUNT", "PENDING PDC", "BOUNCED AMOUNT", "COMBINED VALUE"]
        for col_idx, text in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = DARK_BG

            p = cell.text_frame.paragraphs[0]
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = GOLD
            p.alignment = PP_ALIGN.LEFT if col_idx == 0 else PP_ALIGN.RIGHT

        # Fill Table Rows
        for row_idx, summary in enumerate(property_summaries, start=1):
            row_cells = [
                summary['name'],
                f"AED {summary['cleared']:,.2f}",
                f"AED {summary['pending']:,.2f}",
                f"AED {summary['bounced']:,.2f}",
                f"AED {summary['total']:,.2f}"
            ]

            for col_idx, text in enumerate(row_cells):
                cell = table.cell(row_idx, col_idx)
                cell.text = text

                # Alternate row shading
                cell.fill.solid()
                if row_idx % 2 == 0:
                    cell.fill.fore_color.rgb = RGBColor(248, 250, 252)  # Very light gray
                else:
                    cell.fill.fore_color.rgb = WHITE

                p = cell.text_frame.paragraphs[0]
                p.font.name = "Arial"
                p.font.size = Pt(10.5)
                p.font.color.rgb = TEXT_MAIN

                if col_idx == 4:
                    p.font.bold = True

                p.alignment = PP_ALIGN.LEFT if col_idx == 0 else PP_ALIGN.RIGHT

        # Save to file buffer and return
        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)

        filename = f"Performance_Report_{timezone.now().strftime('%Y%m%d')}.pptx"
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    # -------------------------------------------------------------
    # REST OF VIEW (CSV EXPORT & HTML RENDER)
    # -------------------------------------------------------------
    if export_format == 'csv':
        response = HttpResponse(content_type='text/csv')
        response[
            'Content-Disposition'] = f'attachment; filename="Comprehensive_Report_{timezone.now().strftime("%Y%m%d")}.csv"'
        writer = csv.writer(response)
        writer.writerow(['Property', 'Unit', 'Tenant Name', 'Cheque/Ref Number', 'Amount (AED)', 'Due Date', 'Status'])
        for row in reports_data:
            writer.writerow([
                row.contract.unit.property_asset.name,
                row.contract.unit.unit_number,
                row.contract.tenant.full_name,
                row.cheque_number,
                row.amount,
                row.due_date.strftime("%Y-%m-%d") if row.due_date else 'N/A',
                row.status
            ])
        return response

    properties = Property.objects.filter(owner=landlord)
    notifications = get_dynamic_notifications(landlord)
    total_amount = reports_data.aggregate(Sum('amount'))['amount__sum'] or 0

    return render(request, 'realestate/reports.html', {
        'landlord': landlord,
        'reports_data': reports_data,
        'properties': properties,
        'total_amount': total_amount,
        'filters': {
            'property': int(prop_filter) if prop_filter.isdigit() else '',
            'unit': unit_filter,
            'tenant': tenant_filter,
            'status': status_filter,
            'date_from': date_from,
            'date_to': date_to
        },
        'notifications': notifications,
        'notif_count': len(notifications)
    })
@login_required(login_url='/login/')
def upload_document(request, landlord_id):
    if request.method == 'POST':
        prop_id = request.POST.get('property_id')
        title = request.POST.get('title')
        file = request.FILES.get('document')

        if prop_id and title and file:
            property_obj = get_object_or_404(Property, id=prop_id)
            from .models import PropertyDocument
            PropertyDocument.objects.create(property_asset=property_obj, title=title, file=file)

    return redirect('owner_view_properties', landlord_id=landlord_id)


@login_required(login_url='/login/')
def owner_drive(request, landlord_id):
    landlord = get_object_or_404(Landlord, id=landlord_id)
    if not request.user.is_superuser:
        if not hasattr(request.user, 'landlord_profile') or request.user.landlord_profile.id != landlord.id:
            return HttpResponseForbidden("<h1>403</h1>")

    search_query = request.GET.get('search', '').strip()
    folder_id = request.GET.get('folder')
    current_folder = None

    if search_query:
        folders = DocumentFolder.objects.filter(owner=landlord, name__icontains=search_query)
        files = DocumentFile.objects.filter(owner=landlord, title__icontains=search_query)
    else:
        if folder_id:
            current_folder = get_object_or_404(DocumentFolder, id=folder_id, owner=landlord)
            folders = DocumentFolder.objects.filter(owner=landlord, parent=current_folder).order_by('-created_at')
            files = DocumentFile.objects.filter(owner=landlord, folder=current_folder).order_by('-uploaded_at')
        else:
            folders = DocumentFolder.objects.filter(owner=landlord, parent__isnull=True).order_by('-created_at')
            files = DocumentFile.objects.filter(owner=landlord, folder__isnull=True).order_by('-uploaded_at')

    notifications = get_dynamic_notifications(landlord)

    return render(request, 'realestate/owner_drive.html', {
        'landlord': landlord,
        'folders': folders,
        'files': files,
        'current_folder': current_folder,
        'search_query': search_query,
        'notifications': notifications,
        'notif_count': len(notifications)
    })


@login_required(login_url='/login/')
def drive_create_folder(request, landlord_id):
    if request.method == 'POST':
        name = request.POST.get('folder_name')
        parent_id = request.POST.get('parent_id')
        landlord = get_object_or_404(Landlord, id=landlord_id)
        if name:
            parent_folder = DocumentFolder.objects.get(id=parent_id) if parent_id else None
            DocumentFolder.objects.create(owner=landlord, name=name, parent=parent_folder)

        if parent_id:
            return redirect(f"/owner/{landlord_id}/drive/?folder={parent_id}")
    return redirect('owner_drive', landlord_id=landlord_id)


@login_required(login_url='/login/')
def drive_rename_item(request, landlord_id, item_type, item_id):
    if request.method == 'POST':
        new_name = request.POST.get('new_name')
        if new_name:
            if item_type == 'folder':
                folder = get_object_or_404(DocumentFolder, id=item_id, owner_id=landlord_id)
                folder.name = new_name
                folder.save()
            elif item_type == 'file':
                file_obj = get_object_or_404(DocumentFile, id=item_id, owner_id=landlord_id)
                file_obj.title = new_name
                file_obj.save()
    return redirect(request.META.get('HTTP_REFERER', 'owner_drive'))


@login_required(login_url='/login/')
def drive_delete_item(request, landlord_id, item_type, item_id):
    if request.method == 'POST':
        if item_type == 'folder':
            folder = get_object_or_404(DocumentFolder, id=item_id, owner_id=landlord_id)
            folder.delete()
        elif item_type == 'file':
            file_obj = get_object_or_404(DocumentFile, id=item_id, owner_id=landlord_id)
            file_obj.delete()
    return redirect(request.META.get('HTTP_REFERER', 'owner_drive'))


@login_required(login_url='/login/')
def drive_upload_file(request, landlord_id):
    if request.method == 'POST':
        title = request.POST.get('title')
        file = request.FILES.get('document')
        folder_id = request.POST.get('folder_id')
        landlord = get_object_or_404(Landlord, id=landlord_id)

        if title and file:
            folder = DocumentFolder.objects.get(id=folder_id) if folder_id else None
            DocumentFile.objects.create(owner=landlord, folder=folder, title=title, file=file)

        if folder_id:
            return redirect(f"/owner/{landlord_id}/drive/?folder={folder_id}")

    return redirect('owner_drive', landlord_id=landlord_id)


@login_required(login_url='/login/')
def owner_approvals(request, landlord_id):
    landlord = get_object_or_404(Landlord, id=landlord_id)
    if not request.user.is_superuser:
        if not hasattr(request.user, 'landlord_profile') or request.user.landlord_profile.id != landlord.id:
            return HttpResponseForbidden("<h1>403</h1>")

    docs = ApprovalDocument.objects.filter(
        Q(landlord=landlord) | Q(landlord__isnull=True)
    ).order_by('-doc_date')

    unapproved = docs.filter(status='UnApproved')
    approved = docs.filter(status='Approved')
    rejected = docs.filter(status='Rejected')

    notifications = get_dynamic_notifications(landlord)

    return render(request, 'realestate/approvals.html', {
        'landlord': landlord,
        'unapproved': unapproved,
        'approved': approved,
        'rejected': rejected,
        'notifications': notifications,
        'notif_count': len(notifications)
    })



@login_required(login_url='/login/')
def delete_property_document(request, landlord_id, doc_id):
    if request.method == 'POST':
        doc = get_object_or_404(PropertyDocument, id=doc_id, property_asset__owner__id=landlord_id)
        if doc.file:
            doc.file.delete()
        doc.delete()

    return redirect('owner_view_properties', landlord_id=landlord_id)